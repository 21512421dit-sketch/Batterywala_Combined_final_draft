import calendar
import io
import json
import os
import re
import smtplib
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import formataddr
from pathlib import Path
from html import escape as html_escape
from xml.sax.saxutils import escape as xml_escape

from flask import Blueprint, Response, current_app, jsonify, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from werkzeug.security import generate_password_hash

from . import db
from .models import Delivery, EngineCentre, EngineCentreContact, EngineWhatsAppAdmin, EngineWhatsAppTemplate, Lead, Submission, User


bp = Blueprint('portal', __name__)
CONSENT_VERSION = '2026-09'
SITES = {'batterywala': 'BatteryWala', 'engine_dcarb': 'Engine D-Carb'}
MACHINES = ('DCC AD6000', 'DCC PD6000', 'DCC PD10K', 'DCC PD15K')
ENGINE_CENTRES = (
    ('chikalthana-midc', 'Chikalthana MIDC', 'Care4Earth Enterprises, F-7/3, Chikalthana MIDC, Opp Flemingo Housing Society, Chhatrapati Sambhajinagar, Maharashtra.', '7727005151'),
    ('waluj-midc', 'Waluj MIDC', 'Siddhivinayak Wheel Alignment Center, Opp Tirupati Hospital, Mahaveer Chowk, Waluj MIDC, Chhatrapati Sambhajinagar, Maharashtra.', '8796344838'),
    ('beed-bypass', 'Beed By Pass', 'Master Garage, Opp Kamalnayan Bajaj Hospital, Beside LPG Fuel Pump, Beed By Pass Road, Chhatrapati Sambhajinagar, Maharashtra.', '9923778777'),
    ('garkheda', 'Garkheda', 'Fox 3D Wheel Alignment, Opp Sports Olympia Complex, Sut Girni Chowk, Garkheda, Chhatrapati Sambhajinagar, Maharashtra.', '9518537075'),
)
ENGINE_FAQS = (
    ('What is engine decarbonization?', 'It is a workshop service intended to address carbon deposits that accumulate during engine operation. Engine D-Carb uses a guided HHO cycle while the engine remains assembled.'),
    ('Is engine decarbonization good and safe for an engine?', 'Suitability depends on the vehicle and its condition. Engine D-Carb is designed as a non-invasive, chemical-free process and should be performed by a trained technician after basic vehicle checks.'),
    ('How much does it cost to decarbonize an engine?', 'Cost varies with vehicle class, fuel type, engine capacity and service location. Submit your vehicle details for a tentative quotation; the service team confirms the final price.'),
    ('Does decarbonizing increase mileage?', 'It may help recover efficiency lost to carbon build-up, but results vary with engine condition, maintenance, fuel quality and driving style. No fixed mileage improvement is guaranteed.'),
    ('How long does engine decarbonization take?', 'The Engine D-Carb service typically takes 45 to 90 minutes, depending on the vehicle, inspection and selected cycle.'),
    ('What signs may indicate carbon build-up?', 'Possible signs include rough idling, sluggish response, visible exhaust smoke or a change in Engine Performance. These symptoms can have other causes, so a vehicle check remains important.'),
    ('How often should an engine be decarbonized?', 'There is no universal interval for every vehicle. Driving pattern, symptoms, mileage, maintenance history and the manufacturer’s guidance should inform the decision.'),
)
EXPORT_COLUMNS = {
    'created_at': 'Submitted at', 'form_kind': 'Form type', 'name': 'Customer name',
    'phone': 'Phone', 'email': 'Email', 'employee': 'Submitted by',
    'details': 'Submitted details', 'result': 'Quotation result',
    'consent': 'Consent', 'expires_at': 'Retention expiry',
}
FIELD_LABELS = {
    'area': 'Area', 'businessDetails': 'Business details', 'companyAddress': 'Company address',
    'customerName': 'Customer name', 'engineCc': 'Engine capacity (CC)', 'enquiryType': 'Enquiry type',
    'expertMachine': 'Recommended machine', 'fuelType': 'Fuel type', 'kilometres': 'Kilometres driven',
    'machineCity': 'City', 'machineEmail': 'Email address', 'machinePhone': 'Phone number',
    'machinePin': 'PIN code', 'passingYear': 'Registration year', 'representativeName': 'Representative name',
    'serviceCity': 'City', 'serviceDetails': 'Service details', 'serviceEmail': 'Email address',
    'servicePhone': 'Phone number',
    'servicePin': 'PIN code', 'vehicleBrand': 'Vehicle brand', 'vehicleModel': 'Vehicle model',
    'vehicleType': 'Vehicle type', 'indicative_cost': 'Indicative cost', 'machine': 'Recommended machine',
    'selectedCentre': 'Preferred service centre', 'message': 'Message', 'note': 'Note',
    'number': 'Quotation number', 'status': 'Status',
}
HIDDEN_DISPLAY_FIELDS = {'consent', 'whatsappConsent', 'emailQuote', 'token'}


def field_label(key):
    if key in FIELD_LABELS:
        return FIELD_LABELS[key]
    words = re.sub(r'([a-z0-9])([A-Z])', r'\1 \2', key).replace('_', ' ').strip()
    return words[:1].upper() + words[1:]


def display_value(value, key=''):
    if value is True:
        return 'Yes'
    if value is False:
        return 'No'
    if value in (None, ''):
        return 'Not provided'
    if key == 'indicative_cost':
        return f"₹{int(value):,}"
    if key == 'status':
        return str(value).replace('_', ' ').title()
    if isinstance(value, (list, tuple)):
        return ', '.join(str(item) for item in value) or 'Not provided'
    return str(value)


def display_fields(data, prefix=''):
    fields = []
    for key, value in data.items():
        if key in HIDDEN_DISPLAY_FIELDS:
            continue
        aliases = {'name': ('customerName', 'representativeName'), 'phone': ('servicePhone', 'machinePhone'),
                   'email': ('serviceEmail', 'machineEmail')}
        if key in aliases and any(alias in data for alias in aliases[key]):
            continue
        label = f'{prefix} — {field_label(key)}' if prefix else field_label(key)
        if isinstance(value, dict):
            fields.extend(display_fields(value, label))
        else:
            fields.append({'label': label, 'value': display_value(value, key)})
    return fields


def fields_as_text(fields):
    return '\n'.join(f"{field['label']}: {field['value']}" for field in fields)


def utcnow():
    return datetime.now(timezone.utc)


def retention_expiry(value=None):
    value = value or utcnow()
    year, month = value.year + 2, value.month
    day = min(value.day, calendar.monthrange(year, month)[1])
    return value.replace(year=year, month=month, day=day)


def ensure_engine_centres():
    for index, (key, name, address, default_phone) in enumerate(ENGINE_CENTRES, start=1):
        centre = EngineCentre.query.filter_by(key=key).first()
        if centre:
            if not centre.city:
                centre.city = 'Chhatrapati Sambhajinagar'
            if not centre.state:
                centre.state = 'Maharashtra'
            continue
        centre = EngineCentre(key=key, name=name, address=address, city='Chhatrapati Sambhajinagar',
                              state='Maharashtra', sort_order=index)
        centre.contacts.append(EngineCentreContact(contact_name='Primary centre contact', phone=default_phone))
        db.session.add(centre)
    db.session.commit()


def ensure_engine_whatsapp_admin():
    if db.session.get(EngineWhatsAppAdmin, 1):
        return
    phone = re.sub(r'\D', '', os.getenv('ENGINE_DCARB_ADMIN_WHATSAPP_NUMBER') or '9067671513')
    if len(phone) == 12 and phone.startswith('91'):
        phone = phone[2:]
    if re.fullmatch(r'[0-9]{10}', phone):
        db.session.add(EngineWhatsAppAdmin(id=1, phone=phone))
        db.session.commit()


def centre_dict(centre, include_contacts=False):
    data = {'key': centre.key, 'name': centre.name, 'address': centre.address,
            'city': centre.city, 'state': centre.state}
    if include_contacts:
        data['contacts'] = [{'id': item.id, 'name': item.contact_name, 'phone': item.phone}
                            for item in centre.contacts]
    return data


def service_messages(form, result, centre):
    vehicle = f"{form['vehicleType']} — {form['vehicleBrand']} {form['vehicleModel']}, {form['passingYear']}, {form['fuelType']}, {form['engineCc']} CC, {int(form['kilometres']):,} km"
    numbers = [item.phone for item in centre.contacts] if centre else []
    centre_details = f"{result['centre']['name']}\n{result['centre']['address']}"
    if numbers:
        centre_details += f"\nContact: {', '.join(numbers)}"
    location_label = ('Your requested service area is' if result['centre']['key'] == 'other'
                      else 'Your selected nearest centre details are as follows')
    customer = (
        f"Thank you {form['customerName']} for the enquiry of Engine D-Carb Service.\n\n"
        f"Vehicle details:\n{vehicle}\n\n"
        f"{location_label}:\n{centre_details}"
    )
    owner = (
        f"{form['customerName']} is interested to avail the Engine D-Carb Service. "
        f"The contact number is {form['servicePhone']}.\n\n"
        f"Vehicle details:\n{vehicle}\n\n"
        f"An estimated cost of ₹{result['indicative_cost']:,} is given to the customer for availing Engine D-Carb service.\n\n"
        "Kindly contact the customer and book the appointment."
    )
    sender_number = normalize_whatsapp_number(os.getenv('ENGINE_DCARB_WHATSAPP_NUMBER') or '919607069191')
    return {'customer': customer, 'centre': owner, 'centre_numbers': numbers, 'sender_number': sender_number}


def normalize_whatsapp_number(value):
    """Return a WhatsApp recipient in international digit format."""
    digits = re.sub(r'\D', '', str(value or ''))
    if len(digits) == 10:
        digits = '91' + digits
    return digits


def send_whatsapp_template(target, template_name, parameters, recipient):
    """Send one approved WhatsApp template without exposing credentials to the browser."""
    phone_number_id = (os.getenv('WHATSAPP_PHONE_NUMBER_ID') or '').strip()
    access_token = (os.getenv('WHATSAPP_ACCESS_TOKEN') or '').strip()
    language = (os.getenv('WHATSAPP_TEMPLATE_LANGUAGE') or 'en_IN').strip()
    version = (os.getenv('WHATSAPP_API_VERSION') or 'v23.0').strip()
    normalized_target = normalize_whatsapp_number(target)
    if not phone_number_id or not access_token or not template_name:
        return {'recipient': recipient, 'target': normalized_target, 'status': 'not_configured',
                'detail': 'WhatsApp API configuration is incomplete.'}
    if current_app.config.get('TESTING') and not current_app.config.get('WHATSAPP_ALLOW_TEST_DELIVERY'):
        return {'recipient': recipient, 'target': normalized_target, 'status': 'test_skipped',
                'detail': 'External WhatsApp delivery is disabled during tests.'}
    if not re.fullmatch(r'v\d+\.\d+', version):
        version = 'v23.0'
    payload = {
        'messaging_product': 'whatsapp', 'to': normalized_target, 'type': 'template',
        'template': {
            'name': template_name, 'language': {'code': language},
            'components': [{'type': 'body', 'parameters': [
                {'type': 'text', 'text': str(value)} for value in parameters
            ]}],
        },
    }
    if not parameters:
        payload['template'].pop('components', None)
    graph_request = urllib.request.Request(
        f'https://graph.facebook.com/{version}/{phone_number_id}/messages',
        data=json.dumps(payload, ensure_ascii=False).encode('utf-8'),
        headers={'Authorization': f'Bearer {access_token}', 'Content-Type': 'application/json'},
        method='POST',
    )
    try:
        with urllib.request.urlopen(graph_request, timeout=20) as response:
            response_data = json.loads(response.read().decode('utf-8'))
        message_id = ((response_data.get('messages') or [{}])[0].get('id') or '')
        return {'recipient': recipient, 'target': normalized_target, 'status': 'accepted' if message_id else 'failed',
                'message_id': message_id, 'detail': message_id or 'Meta returned no message ID.'}
    except urllib.error.HTTPError as error:
        try:
            error_data = json.loads(error.read().decode('utf-8'))
            meta_error = error_data.get('error') or {}
            detail = (meta_error.get('error_data') or {}).get('details') or meta_error.get('message') or f'Meta returned HTTP {error.code}.'
        except (ValueError, UnicodeDecodeError):
            detail = f'Meta returned HTTP {error.code}.'
        current_app.logger.warning('WhatsApp template delivery failed for %s: %s', recipient, detail)
        return {'recipient': recipient, 'target': normalized_target, 'status': 'failed', 'detail': detail[:500]}
    except (OSError, ValueError) as error:
        current_app.logger.warning('WhatsApp template delivery failed for %s: %s', recipient, error)
        return {'recipient': recipient, 'target': normalized_target, 'status': 'failed',
                'detail': 'Unable to reach the WhatsApp service.'}


def approved_whatsapp_template(name):
    """Switch to a replacement only after Meta reports it approved."""
    account_id = (os.getenv('WHATSAPP_BUSINESS_ACCOUNT_ID') or '').strip()
    token = (os.getenv('WHATSAPP_ACCESS_TOKEN') or '').strip()
    if not account_id or not token or not re.fullmatch(r'\d+', account_id):
        return False
    version = (os.getenv('WHATSAPP_API_VERSION') or 'v23.0').strip()
    if not re.fullmatch(r'v\d+\.\d+', version):
        version = 'v23.0'
    query = urllib.parse.urlencode({'name': name, 'fields': 'name,status,language', 'limit': 10})
    lookup = urllib.request.Request(
        f'https://graph.facebook.com/{version}/{account_id}/message_templates?{query}',
        headers={'Authorization': f'Bearer {token}'})
    try:
        with urllib.request.urlopen(lookup, timeout=5) as response:
            templates = json.load(response).get('data', [])
    except (OSError, ValueError):
        return False
    language = (os.getenv('WHATSAPP_TEMPLATE_LANGUAGE') or 'en_IN').strip()
    return any(item.get('name') == name and item.get('language') == language
               and item.get('status') == 'APPROVED' for item in templates)


def approved_replacement_templates(names):
    """Read replacement approvals in one request; keep approved originals on lookup failure."""
    account_id = (os.getenv('WHATSAPP_BUSINESS_ACCOUNT_ID') or '').strip()
    token = (os.getenv('WHATSAPP_ACCESS_TOKEN') or '').strip()
    if not account_id or not token or not re.fullmatch(r'\d+', account_id):
        return set()
    version = (os.getenv('WHATSAPP_API_VERSION') or 'v23.0').strip()
    if not re.fullmatch(r'v\d+\.\d+', version):
        version = 'v23.0'
    query = urllib.parse.urlencode({'fields': 'name,status,language', 'limit': 100})
    lookup = urllib.request.Request(
        f'https://graph.facebook.com/{version}/{account_id}/message_templates?{query}',
        headers={'Authorization': f'Bearer {token}'})
    try:
        with urllib.request.urlopen(lookup, timeout=5) as response:
            templates = json.load(response).get('data', [])
    except (OSError, ValueError):
        return set()
    language = (os.getenv('WHATSAPP_TEMPLATE_LANGUAGE') or 'en_IN').strip()
    return {item['name'] for item in templates if item.get('name') in names
            and item.get('language') == language and item.get('status') == 'APPROVED'}


def engine_admin_number():
    configured = db.session.get(EngineWhatsAppAdmin, 1)
    return normalize_whatsapp_number(
        configured.phone if configured else os.getenv('ENGINE_DCARB_ADMIN_WHATSAPP_NUMBER') or '919067671513')


def india_submission_time():
    current = datetime.now(timezone(timedelta(hours=5, minutes=30)))
    hour = current.strftime('%I').lstrip('0') or '0'
    return f"{current.day} {current.strftime('%b %Y')}, {hour}:{current.strftime('%M %p')} IST"


def machine_template_values(form):
    return {
        'representative_name': form['representativeName'],
        'machine_phone': form['machinePhone'], 'machine_email': form['machineEmail'],
        'company_address': form['companyAddress'], 'business_details': form['businessDetails'],
        'machine_city': form['machineCity'], 'machine_pin': form['machinePin'],
        'machine_details': form.get('machineDetails') or 'Not provided',
        'submitted_at': india_submission_time(),
    }


def centre_head_phone(centre):
    if not centre or not centre.contacts:
        return 'To be confirmed'
    named_head = next((item for item in centre.contacts if 'head' in item.contact_name.lower()), None)
    return (named_head or centre.contacts[0]).phone


def send_unique_whatsapp(planned):
    """Send one message per phone, keeping the most actionable role's template."""
    role_priority = {'customer': 0, 'centre': 1, 'admin': 2}
    selected = {}
    for index, (recipient, target, name, parameters) in enumerate(planned):
        number = normalize_whatsapp_number(target)
        role = recipient.split(':', 1)[0]
        priority = role_priority[role]
        if number not in selected or priority > selected[number][0]:
            selected[number] = (priority, index, recipient, number, name, parameters)
    return [send_whatsapp_template(number, name, parameters, recipient)
            for _, _, recipient, number, name, parameters in sorted(selected.values(), key=lambda item: item[1])]


def send_engine_whatsapp(form, result, centre=None):
    from .engine_templates import render_template_parameters, render_spaced_parameters
    if os.getenv('WHATSAPP_TEST_MODE', '').lower() == 'true':
        result['whatsapp_test_mode'] = True
        customer_phone = form.get('servicePhone') or form.get('machinePhone')
        return [send_whatsapp_template(customer_phone, 'hello_world', [], 'customer')]

    if form['enquiryType'] == 'machine':
        values = machine_template_values(form)
        customer_template, customer_parameters = render_template_parameters('machine_customer', values)
        replacement = 'engine_dcarb_machine_confirmation_v2'
        approved = approved_replacement_templates({replacement, 'engine_dcarb_admin_centre_lead_v3'})
        if replacement in approved and not db.session.get(EngineWhatsAppTemplate, 'machine_customer'):
            customer_template, customer_parameters = render_spaced_parameters(
                replacement, ['representative_name', 'company_address'], values)
        spaced_template, spaced_parameters = render_template_parameters('machine_admin_spaced', values)
        replacement = 'engine_dcarb_admin_centre_lead_v3'
        if db.session.get(EngineWhatsAppTemplate, 'machine_admin'):
            admin_template, admin_parameters = render_template_parameters('machine_admin', values)
        elif db.session.get(EngineWhatsAppTemplate, 'machine_admin_spaced') and approved_whatsapp_template(spaced_template):
            admin_template, admin_parameters = spaced_template, spaced_parameters
        elif replacement in approved:
            admin_template, admin_parameters = render_spaced_parameters(
                replacement, ['representative_name', 'machine_phone', 'machine_email',
                              'company_address', 'business_details', 'machine_city',
                              'machine_pin', 'machine_details', 'submitted_at'], values)
        elif approved_whatsapp_template(spaced_template):
            admin_template, admin_parameters = spaced_template, spaced_parameters
        else:
            admin_template, admin_parameters = render_template_parameters('machine_admin', values)
        return send_unique_whatsapp([
            ('customer', form['machinePhone'], customer_template, customer_parameters),
            ('admin', engine_admin_number(), admin_template, admin_parameters),
        ])

    vehicle = (f"{form['vehicleType']} — {form['vehicleBrand']} {form['vehicleModel']} "
               f"({form['passingYear']}), {form['fuelType']}, {form['engineCc']} CC")
    centre_heads = ', '.join(
        f'{contact.contact_name} ({contact.phone})' for contact in centre.contacts
    ) if centre and centre.contacts else 'To be assigned'
    values = {
        'customer_name': form['customerName'], 'customer_phone': form['servicePhone'],
        'vehicle': vehicle, 'kilometres': f"{int(form['kilometres']):,} km",
        'service_details': form.get('serviceDetails') or 'Not provided',
        'cost': result['indicative_cost'], 'centre_name': result['centre']['name'],
        'centre_address': result['centre']['address'],
        'centre_phone': centre_head_phone(centre),
        'centre_heads': centre_heads, 'submitted_at': india_submission_time(),
    }
    customer_template, customer_parameters = render_template_parameters('service_customer', values)
    centre_template, centre_parameters = render_template_parameters('service_centre', values)
    admin_template, admin_parameters = render_template_parameters('service_admin', values)
    replacements = (
        ('engine_dcarb_service_quote_v2', ['customer_name', 'vehicle', 'cost',
                                           'centre_name', 'centre_address', 'centre_phone']),
        ('engine_dcarb_service_centre_lead_v2', ['customer_name', 'customer_phone',
                                                 'vehicle', 'cost', 'centre_name', 'submitted_at']),
        ('engine_dcarb_service_admin_lead_v2', ['customer_name', 'customer_phone',
                                                'vehicle', 'kilometres', 'service_details',
                                                'cost', 'centre_name', 'centre_address',
                                                'centre_heads', 'submitted_at']),
    )
    approved = approved_replacement_templates({name for name, _ in replacements})
    if replacements[0][0] in approved and not db.session.get(EngineWhatsAppTemplate, 'service_customer'):
        customer_template, customer_parameters = render_spaced_parameters(*replacements[0], values)
    if replacements[1][0] in approved and not db.session.get(EngineWhatsAppTemplate, 'service_centre'):
        centre_template, centre_parameters = render_spaced_parameters(*replacements[1], values)
    if replacements[2][0] in approved and not db.session.get(EngineWhatsAppTemplate, 'service_admin'):
        admin_template, admin_parameters = render_spaced_parameters(*replacements[2], values)
    planned = [('customer', form['servicePhone'], customer_template, customer_parameters)]
    if centre:
        for contact in centre.contacts:
            planned.append((f'centre:{contact.contact_name}', contact.phone,
                            centre_template, centre_parameters))
    planned.append(('admin', engine_admin_number(), admin_template, admin_parameters))
    return send_unique_whatsapp(planned)


def send_engine_emails(form, result, customer_email):
    """Send the customer confirmation and an internal, replyable enquiry copy."""
    sender = (os.getenv('SMTP_FROM') or os.getenv('SMTP_USERNAME') or '').strip()
    engine_email = (os.getenv('ENGINE_DCARB_EMAIL') or sender).strip()
    if not os.getenv('SMTP_HOST') or not sender or not engine_email:
        return [
            {'recipient': 'customer', 'status': 'not_configured'},
            {'recipient': 'engine_dcarb', 'status': 'not_configured'},
        ]

    kind = form['enquiryType']
    customer_name = form.get('customerName') or form.get('representativeName') or 'Customer'
    if kind == 'service':
        centre = result['centre']
        customer_subject = 'Your Engine D-Carb service quotation'
        customer_body = (
            f"Hello {customer_name},\n\n"
            f"Thank you for your Engine D-Carb service enquiry.\n\n"
            f"Indicative service cost: ₹{result['indicative_cost']:,}\n"
            f"Selected centre: {centre['name']}\n{centre['address']}\n\n"
            f"Vehicle: {form['vehicleBrand']} {form['vehicleModel']} ({form['passingYear']})\n"
            f"Fuel: {form['fuelType']}\nEngine capacity: {form['engineCc']} CC\n"
            f"Kilometres: {int(form['kilometres']):,} km\n\n"
            "This is an indicative quotation. The service team will confirm the final price and appointment."
        )
        internal_subject = f"New Engine D-Carb service enquiry — {customer_name}"
    else:
        customer_subject = 'Your Engine D-Carb machine enquiry'
        customer_body = (
            f"Hello {customer_name},\n\n"
            "Thank you for your Engine D-Carb machine enquiry. Our team will review your business "
            "requirements and contact you with the suitable configuration and quotation.\n\n"
            "Regards,\nEngine D-Carb"
        )
        internal_subject = f"New Engine D-Carb machine enquiry — {customer_name}"

    internal_body = (
        "A customer submitted an Engine D-Carb enquiry and requested email follow-up.\n"
        "Reply to this email to contact the customer directly.\n\n"
        + fields_as_text(display_fields(form))
        + "\n\nResult:\n"
        + fields_as_text(display_fields(result))
    )
    messages = []
    for recipient, target, subject, body, reply_to in (
        ('customer', customer_email, customer_subject, customer_body, engine_email),
        ('engine_dcarb', engine_email, internal_subject, internal_body, customer_email),
    ):
        message = EmailMessage()
        message['Subject'] = subject
        message['From'] = formataddr(('Engine D-Carb', sender))
        message['To'] = target
        message['Reply-To'] = reply_to
        message.set_content(body)
        messages.append((recipient, message))

    deliveries = []
    try:
        with smtplib.SMTP(os.environ['SMTP_HOST'], int(os.getenv('SMTP_PORT', '587')), timeout=15) as smtp:
            if os.getenv('SMTP_USE_TLS', 'true').lower() == 'true':
                smtp.starttls()
            if os.getenv('SMTP_USERNAME'):
                smtp.login(os.environ['SMTP_USERNAME'], os.getenv('SMTP_PASSWORD'))
            for recipient, message in messages:
                smtp.send_message(message)
                deliveries.append({'recipient': recipient, 'status': 'sent'})
    except Exception:
        sent = {item['recipient'] for item in deliveries}
        deliveries.extend({'recipient': recipient, 'status': 'failed'} for recipient, _ in messages if recipient not in sent)
    return deliveries


def purge_expired_submissions():
    now = utcnow()
    expired = Submission.query.filter(Submission.expires_at <= now).all()
    if not expired:
        return 0
    for item in expired:
        delete_submission_data(item)
    db.session.commit()
    return len(expired)


def delete_submission_data(item):
    """Delete one request and its owned lead, deliveries, and WhatsApp status records."""
    from .models import WhatsAppMessage
    messages = WhatsAppMessage.query.filter_by(submission_id=item.id).all()
    delivery_ids = {message.delivery_id for message in messages if message.delivery_id}
    for message in messages:
        db.session.delete(message)
    db.session.flush()
    if delivery_ids:
        Delivery.query.filter(Delivery.id.in_(delivery_ids)).delete(synchronize_session=False)
    Delivery.query.filter_by(submission_id=item.id).delete(synchronize_session=False)
    if item.lead_id:
        Delivery.query.filter_by(lead_id=item.lead_id).delete(synchronize_session=False)
        lead = db.session.get(Lead, item.lead_id)
        if lead:
            db.session.delete(lead)
    db.session.delete(item)


def record_submission(site, form_kind, form, result=None, lead_id=None, consented=True):
    if not consented:
        return None
    actor_id = current_user.id if current_user.is_authenticated and not current_user.is_admin else None
    submission = Submission(
        site=site, form_kind=form_kind, name=form.get('name') or form.get('customerName') or form.get('representativeName'),
        email=form.get('email') or form.get('machineEmail'), phone=form.get('phone') or form.get('servicePhone') or form.get('machinePhone'),
        form_json=json.dumps(form, ensure_ascii=False), result_json=json.dumps(result or {}, ensure_ascii=False),
        consented=True, consent_version=CONSENT_VERSION, submitted_by_id=actor_id,
        lead_id=lead_id, expires_at=retention_expiry(),
    )
    db.session.add(submission)
    return submission


def admin_required(fn):
    from functools import wraps

    @wraps(fn)
    @login_required
    def inner(*args, **kwargs):
        if not current_user.is_admin:
            return ('Forbidden', 403)
        return fn(*args, **kwargs)

    return inner


def engine_public_origin():
    domain = (os.getenv('ENGINE_DCARB_DOMAIN') or '').strip().split(':')[0].lower()
    if domain:
        return f'https://{domain}'
    forwarded_proto = request.headers.get('X-Forwarded-Proto', '').split(',')[0].strip().lower()
    scheme = 'https' if forwarded_proto == 'https' or request.host.lower().endswith('.ngrok-free.dev') else request.scheme
    return f'{scheme}://{request.host}'


def engine_canonical_url():
    return engine_public_origin() + ('/' if os.getenv('ENGINE_DCARB_DOMAIN', '').strip() else request.path)


def engine_site_response():
    source = Path(current_app.root_path).parent / 'docs' / 'engine_dcarb_latest.html'
    html = source.read_text(encoding='utf-8')
    html = re.sub(r'<script type="application/ld\+json">.*?</script>', '', html, count=1, flags=re.DOTALL)
    title = 'Engine D-Carb | Engine Decarbonization Service & Machines'
    description = ('Explore Engine D-Carb vehicle decarbonization services and HHO machines for workshops. '
                   'Choose a service centre or enquire about a machine in India.')
    html = re.sub(r'<title>.*?</title>', f'<title>{title}</title>', html, count=1, flags=re.DOTALL)
    html = re.sub(r'<meta name="description" content="[^"]*">',
                  f'<meta name="description" content="{html_escape(description, quote=True)}">', html, count=1)
    html = re.sub(r'\s*<meta name="keywords" content="[^"]*">', '', html, count=1)
    public_email = 'enquiry@enginedcarb.com'
    html = html.replace('rahul.care4earth@outlook.com', public_email)
    html = html.replace('Your information stays in your browser until you choose to send it.',
                        'Your information is stored for 24 months only after you accept the consent notice and submit.')
    html = html.replace('Submitting prepares the summary on this device. Nothing is uploaded automatically.',
                        'Submitting stores this request only after consent, then prepares the summary on this device.')
    html = html.replace('Complete the form, then continue on WhatsApp for tentative servicing cost and the nearest service-centre contact details.',
                        'Complete the form, choose your nearest centre and send your enquiry to the service team.')
    html = html.replace('Continue on WhatsApp to receive tentative cost and nearby service-centre details.',
                        'After submitting, open WhatsApp with your enquiry already filled in and press Send.')
    html = html.replace('prints a dated emission report', 'generates an emission report')
    faq_markup = ''.join(f'<details><summary>{question}</summary><p>{answer}</p></details>'
                         for question, answer in ENGINE_FAQS)
    html = re.sub(r'(<div class="faq-list reveal">).*?(</div></div></section>`;)',
                  rf'\1{faq_markup}\2', html, count=1, flags=re.DOTALL)
    legal_links = (
        '<nav class="engine-legal-links" aria-label="Legal information">'
        '<a href="/engine-d-carb/privacy-policy" data-engine-legal="privacy">Privacy Policy</a>'
        '<a href="/engine-d-carb/terms-and-conditions" data-engine-legal="terms">Terms and Conditions</a>'
        '</nav>'
    )
    html = html.replace('<span>Made in India</span></div></footer>',
                        f'<span>Made in India</span>{legal_links}</div></footer>', 1)
    legal_dialog = (
        '<dialog class="engine-legal-dialog" id="engineLegalDialog" aria-labelledby="engineLegalTitle">'
        '<div class="engine-legal-dialog-frame">'
        '<header class="engine-legal-dialog-header"><h2 id="engineLegalTitle" tabindex="-1"></h2>'
        '<button type="button" class="engine-legal-close" aria-label="Close legal information">&times;</button></header>'
        '<div class="engine-legal-dialog-body">'
        '<article class="engine-legal-document" data-engine-legal-content="privacy" hidden>'
        + render_template('engine_legal_content.html', legal_kind='privacy') + '</article>'
        '<article class="engine-legal-document" data-engine-legal-content="terms" hidden>'
        + render_template('engine_legal_content.html', legal_kind='terms') + '</article>'
        '</div></div></dialog>'
    )
    html = html.replace('</body>', legal_dialog + '</body>', 1)
    faq_schema = {'@context': 'https://schema.org', '@type': 'FAQPage', 'mainEntity': [
        {'@type': 'Question', 'name': question,
         'acceptedAnswer': {'@type': 'Answer', 'text': answer}}
        for question, answer in ENGINE_FAQS]}
    canonical = engine_canonical_url()
    origin = engine_public_origin()
    organization_schema = {
        '@context': 'https://schema.org', '@type': 'Organization',
        'name': 'Engine D-Carb', 'url': canonical,
        'email': public_email, 'telephone': '+91 9607069191',
    }
    image_url = origin + '/static/engine-assets/dcarb-technician-connection.webp'
    metadata = (
        f'<link rel="canonical" href="{html_escape(canonical, quote=True)}">'
        '<meta name="robots" content="index,follow">'
        '<meta property="og:type" content="website">'
        f'<meta property="og:title" content="{html_escape(title, quote=True)}">'
        f'<meta property="og:description" content="{html_escape(description, quote=True)}">'
        f'<meta property="og:url" content="{html_escape(canonical, quote=True)}">'
        f'<meta property="og:image" content="{html_escape(image_url, quote=True)}">'
        '<meta name="twitter:card" content="summary_large_image">'
        f'<meta name="twitter:title" content="{html_escape(title, quote=True)}">'
        f'<meta name="twitter:description" content="{html_escape(description, quote=True)}">'
        f'<meta name="twitter:image" content="{html_escape(image_url, quote=True)}">'
    )
    integration = (
        '<link rel="icon" href="/static/engine-assets/engine-dcarb-logo.webp">'
        '<link rel="stylesheet" href="/static/engine-integration.css?v=20260925-1">'
        '<link rel="stylesheet" href="/static/engine-legal.css?v=20260924-1">'
        '<meta name="application-name" content="Engine D-Carb">'
        + metadata
        + ''.join(f'<script type="application/ld+json">{json.dumps(schema, ensure_ascii=False).replace("<", "\\u003c")}</script>'
                  for schema in (organization_schema, faq_schema)) + '</head>'
    )
    scripts = ('<script src="/static/engine-integration.js?v=20260925-1"></script>'
               '<script src="/static/engine-legal.js?v=20260924-1"></script></body>')
    return Response(html.replace('</head>', integration).replace('</body>', scripts), mimetype='text/html')


@bp.get('/engine-d-carb')
def engine_site():
    return engine_site_response()


@bp.get('/sitemap.xml')
def engine_sitemap():
    domain = (os.getenv('ENGINE_DCARB_DOMAIN') or '').strip().split(':')[0].lower()
    if domain and request.host.split(':')[0].lower() != domain:
        return Response('Not found', status=404)
    origin = engine_public_origin()
    home = origin + ('/' if domain else '/engine-d-carb')
    urls = (home, origin + '/engine-d-carb/privacy-policy',
            origin + '/engine-d-carb/terms-and-conditions')
    items = ''.join(f'<url><loc>{xml_escape(url)}</loc></url>' for url in urls)
    return Response('<?xml version="1.0" encoding="UTF-8"?>'
                    '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">'
                    + items + '</urlset>', mimetype='application/xml')


@bp.get('/robots.txt')
def engine_robots():
    domain = (os.getenv('ENGINE_DCARB_DOMAIN') or '').strip().split(':')[0].lower()
    rules = 'User-agent: *\nDisallow: /admin/\nDisallow: /portal\nDisallow: /api/\n'
    if not domain or request.host.split(':')[0].lower() == domain:
        rules += f'Sitemap: {engine_public_origin()}/sitemap.xml\n'
    return Response(rules, mimetype='text/plain')


@bp.get('/engine-d-carb/privacy-policy')
@bp.get('/engine-d-carb/terms-and-conditions')
def engine_legal_page():
    legal_kind = 'privacy' if request.path.endswith('/privacy-policy') else 'terms'
    return render_template('engine_legal_page.html', legal_kind=legal_kind,
                           canonical_url=engine_public_origin() + request.path)


@bp.get('/api/engine-d-carb/centres')
def engine_centres():
    centres = EngineCentre.query.order_by(EngineCentre.sort_order, EngineCentre.id).all()
    return jsonify(centres=[centre_dict(item) for item in centres])


@bp.get('/portal')
@login_required
def workspace():
    if current_user.is_admin:
        return redirect(url_for('main.admin'))
    return render_template('employee.html', machines=MACHINES, consent_version=CONSENT_VERSION)


def valid_engine_payload(payload):
    if not isinstance(payload, dict) or len(payload) > 35:
        return None, 'Submit valid form data.'
    if any(not isinstance(value, (str, bool)) for value in payload.values()):
        return None, 'Submit valid form fields.'
    clean = {key: value.strip() if isinstance(value, str) else value for key, value in payload.items()}
    if any(len(value) > (2000 if key in {'serviceDetails', 'businessDetails', 'machineDetails', 'companyAddress'} else 255)
           for key, value in clean.items() if isinstance(value, str)):
        return None, 'One or more fields are too long.'
    return clean, None


@bp.post('/api/engine-d-carb/quotations')
def engine_quotation():
    form, error = valid_engine_payload(request.get_json(silent=True))
    if error:
        return jsonify(error=error), 400
    consented = form.get('consent') is True or form.get('consent') == 'true'
    if not consented:
        return jsonify(error='Consent is required before we can store your details and prepare the quotation.'), 400
    kind = form.get('enquiryType')
    if kind not in {'service', 'machine'}:
        return jsonify(error='Choose vehicle service or new machine enquiry.'), 400
    employee = current_user.is_authenticated and not current_user.is_admin
    selected_machine = form.get('expertMachine', '') if employee else ''
    if selected_machine and selected_machine not in MACHINES:
        return jsonify(error='Choose a valid Engine D-Carb machine.'), 400
    if kind == 'service':
        required = ('customerName', 'servicePhone', 'vehicleType', 'vehicleBrand', 'vehicleModel', 'passingYear',
                    'fuelType', 'engineCc', 'kilometres', 'selectedCentre')
        if any(not form.get(key) for key in required):
            return jsonify(error='Complete all required vehicle service fields.'), 400
        if not re.fullmatch(r'[0-9]{10}', form['servicePhone']):
            return jsonify(error='Enter a valid 10-digit phone number.'), 400
        try:
            cc, year, kilometres = int(form['engineCc']), int(form['passingYear']), int(form['kilometres'])
        except ValueError:
            return jsonify(error='Enter valid numbers for year, engine CC and kilometres.'), 400
        if not 999 <= cc <= 15000 or not 2010 <= year <= 2027 or kilometres < 0:
            return jsonify(error='Use 999–15,000 CC, a registration year from 2010–2027 and valid kilometres.'), 400
        factor = 1.1 if form['vehicleType'] in {'Car', 'Car/SUV'} else 1.5
        if form['vehicleType'] in {'Car', 'Car/SUV'} and form['fuelType'] == 'Diesel':
            factor = 1.25
        centre = None
        if form['selectedCentre'] == 'other':
            location_required = ('area', 'serviceCity', 'servicePin')
            if any(not form.get(key) for key in location_required):
                return jsonify(error='Enter your area, city and PIN code so we can find the nearest centre.'), 400
            if not re.fullmatch(r'[0-9]{6}', form['servicePin']):
                return jsonify(error='Enter a valid 6-digit PIN code.'), 400
            centre_data = {
                'key': 'other',
                'name': 'Nearest centre to be confirmed',
                'address': f"{form['area']}, {form['serviceCity']} — {form['servicePin']}",
            }
        else:
            centre = EngineCentre.query.filter_by(key=form['selectedCentre']).first()
            if not centre:
                return jsonify(error='Choose a valid Engine D-Carb service centre.'), 400
            centre_data = centre_dict(centre)
        result = {'status': 'ready', 'indicative_cost': round(cc * factor * 1.1),
                  'machine': selected_machine or None, 'centre': centre_data}
        result['messages'] = service_messages(form, result, centre)
        name, phone, email = form['customerName'], form['servicePhone'], form.get('serviceEmail', '')
    else:
        required = ('companyAddress', 'machineEmail', 'representativeName', 'machinePhone', 'businessDetails', 'machineCity', 'machinePin')
        if any(not form.get(key) for key in required):
            return jsonify(error='Complete all required machine enquiry fields.'), 400
        if not re.fullmatch(r'[0-9]{10}', form['machinePhone']) or not re.fullmatch(r'[0-9]{6}', form['machinePin']):
            return jsonify(error='Enter a valid 10-digit phone number and 6-digit PIN code.'), 400
        if not re.fullmatch(r'[^\s@<>]+@[^\s@<>]+\.[^\s@<>]+', form['machineEmail']):
            return jsonify(error='Enter a valid email address.'), 400
        result = {'status': 'received', 'machine': selected_machine or None}
        name, phone, email = form['representativeName'], form['machinePhone'], form['machineEmail']
    email_requested = form.get('emailQuote') is True or form.get('emailQuote') == 'true'
    if email and not re.fullmatch(r'[^\s@<>]+@[^\s@<>]+\.[^\s@<>]+', email):
        return jsonify(error='Enter a valid email address.'), 400
    if email_requested and not email:
        return jsonify(error='Enter a valid email address to receive the quotation.'), 400
    result['whatsapp_delivery'] = send_engine_whatsapp(form, result, centre if kind == 'service' else None)
    result['business_whatsapp_number'] = normalize_whatsapp_number(
        os.getenv('ENGINE_DCARB_WHATSAPP_NUMBER') or '919607069191')
    normalized = dict(form, name=name, phone=phone, email=email)
    submission = record_submission('engine_dcarb', kind, normalized, result, consented=True)
    db.session.flush()
    from .whatsapp_webhook import track_message
    for item in result.get('whatsapp_delivery', []):
        delivery = Delivery(submission_id=submission.id, channel='whatsapp', target=item['target'], status=item['status'],
                            detail=f"{item['recipient']}: {item['detail']}")
        db.session.add(delivery)
        db.session.flush()
        if item.get('message_id'):
            track_message(item['message_id'], delivery, submission)
    db.session.commit()
    if email_requested:
        result['email_delivery'] = send_engine_emails(form, result, email)
    return jsonify(result)


@bp.post('/admin/engine-centres/<int:centre_id>/contacts')
@admin_required
def add_engine_centre_contact(centre_id):
    from .routes import valid_csrf

    if not valid_csrf():
        return ('Invalid CSRF', 400)
    centre = db.get_or_404(EngineCentre, centre_id)
    contact_name = request.form.get('contact_name', '').strip() or 'Centre contact'
    phone = re.sub(r'\D', '', request.form.get('phone', ''))
    if len(contact_name) > 120 or not re.fullmatch(r'[0-9]{10}', phone):
        return redirect(url_for('main.admin', site='engine_dcarb', centre_error='Enter a contact name and valid 10-digit number.'))
    if EngineCentreContact.query.filter_by(centre_id=centre.id, phone=phone).first():
        return redirect(url_for('main.admin', site='engine_dcarb', centre_error='That number is already assigned to this centre.'))
    db.session.add(EngineCentreContact(centre=centre, contact_name=contact_name, phone=phone))
    db.session.commit()
    return redirect(url_for('main.admin', site='engine_dcarb', centre_added='1'))


def centre_fields():
    fields = {key: request.form.get(key, '').strip() for key in ('name', 'address', 'city', 'state')}
    if any(not value for value in fields.values()) or any(len(value) > limit for value, limit in
            ((fields['name'], 120), (fields['address'], 500), (fields['city'], 120), (fields['state'], 120))):
        return None
    return fields


@bp.post('/admin/engine-centres')
@admin_required
def add_engine_centre():
    from uuid import uuid4
    from .routes import valid_csrf
    if not valid_csrf():
        return ('Invalid CSRF', 400)
    fields = centre_fields()
    if not fields:
        return redirect(url_for('main.admin', site='engine_dcarb', centre_error='Enter a name, address, city and state.'))
    key = (re.sub(r'[^a-z0-9]+', '-', fields['name'].lower()).strip('-')[:30] or 'centre') + '-' + uuid4().hex[:8]
    next_order = (db.session.query(db.func.max(EngineCentre.sort_order)).scalar() or 0) + 1
    db.session.add(EngineCentre(key=key, sort_order=next_order, **fields))
    db.session.commit()
    return redirect(url_for('main.admin', site='engine_dcarb', centre_created='1') + '#engine-centres')


@bp.post('/admin/engine-centres/<int:centre_id>/edit')
@admin_required
def edit_engine_centre(centre_id):
    from .routes import valid_csrf
    if not valid_csrf():
        return ('Invalid CSRF', 400)
    centre = db.get_or_404(EngineCentre, centre_id)
    fields = centre_fields()
    if not fields:
        return redirect(url_for('main.admin', site='engine_dcarb', centre_error='Enter a name, address, city and state.'))
    for key, value in fields.items():
        setattr(centre, key, value)
    db.session.commit()
    return redirect(url_for('main.admin', site='engine_dcarb', centre_updated='1') + '#engine-centres')


@bp.post('/admin/engine-centre-contacts/<int:contact_id>/edit')
@admin_required
def edit_engine_centre_contact(contact_id):
    from .routes import valid_csrf
    if not valid_csrf():
        return ('Invalid CSRF', 400)
    contact = db.get_or_404(EngineCentreContact, contact_id)
    name = request.form.get('contact_name', '').strip()
    phone = request.form.get('phone', '').strip()
    if not name or len(name) > 120 or not re.fullmatch(r'[0-9]{10}', phone):
        return redirect(url_for('main.admin', site='engine_dcarb', centre_error='Enter a contact name and valid 10-digit number.'))
    existing = EngineCentreContact.query.filter_by(centre_id=contact.centre_id, phone=phone).first()
    if existing and existing.id != contact.id:
        return redirect(url_for('main.admin', site='engine_dcarb', centre_error='That number is already assigned to this centre.'))
    contact.contact_name, contact.phone = name, phone
    db.session.commit()
    return redirect(url_for('main.admin', site='engine_dcarb', centre_updated='1') + '#engine-centres')


@bp.post('/admin/engine-centre-contacts/<int:contact_id>/delete')
@admin_required
def delete_engine_centre_contact(contact_id):
    from .routes import valid_csrf

    if not valid_csrf():
        return ('Invalid CSRF', 400)
    contact = db.session.get(EngineCentreContact, contact_id)
    if contact:
        db.session.delete(contact)
        db.session.commit()
    return redirect(url_for('main.admin', site='engine_dcarb', centre_deleted='1'))


@bp.post('/admin/employees')
@admin_required
def add_employee():
    from .routes import valid_csrf

    if not valid_csrf():
        return ('Invalid CSRF', 400)
    email = request.form.get('email', '').strip().lower()
    password = request.form.get('password', '')
    if not re.fullmatch(r'[^\s@<>]+@[^\s@<>]+\.[^\s@<>]+', email) or len(password) < 8:
        return redirect(url_for('main.admin', employee_error='Use a valid email and a password of at least 8 characters.'))
    if User.query.filter_by(email=email).first():
        return redirect(url_for('main.admin', employee_error='That email already has an account.'))
    db.session.add(User(email=email, password_hash=generate_password_hash(password), is_admin=False))
    db.session.commit()
    return redirect(url_for('main.admin', employee_added='1'))


def filtered_submissions(args):
    site = args.get('site', 'batterywala')
    if site not in SITES:
        site = 'batterywala'
    query = Submission.query.filter_by(site=site, consented=True)
    kind = args.get('kind', '').strip()
    if kind:
        query = query.filter_by(form_kind=kind)
    employee_id = args.get('employee_id', '').strip()
    if employee_id.isdigit():
        query = query.filter_by(submitted_by_id=int(employee_id))
    search = args.get('search', '').strip()[:120]
    if search:
        from sqlalchemy import or_
        escaped = search.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_')
        pattern = f'%{escaped}%'
        query = query.filter(or_(Submission.name.ilike(pattern, escape='\\'),
                                 Submission.phone.ilike(pattern, escape='\\'),
                                 Submission.email.ilike(pattern, escape='\\'),
                                 Submission.form_json.ilike(pattern, escape='\\')))
    for key, operator in (('from', 'from'), ('to', 'to')):
        raw = args.get(key, '')
        if not raw:
            continue
        try:
            value = datetime.fromisoformat(raw).replace(tzinfo=timezone.utc)
            if operator == 'to':
                value = value.replace(hour=23, minute=59, second=59)
            query = query.filter(Submission.created_at >= value) if operator == 'from' else query.filter(Submission.created_at <= value)
        except ValueError:
            pass
    return site, query.order_by(Submission.created_at.desc())


@bp.post('/admin/submissions/<int:submission_id>/delete')
@admin_required
def delete_submission(submission_id):
    from .routes import valid_csrf
    if not valid_csrf():
        return ('Invalid CSRF', 400)
    item = db.get_or_404(Submission, submission_id)
    site = request.form.get('site', '')
    if site != item.site:
        return ('Invalid site', 400)
    delete_submission_data(item)
    db.session.commit()
    filters = {key: request.form.get(key, '')[:120] for key in ('search', 'from', 'to', 'kind', 'employee_id')}
    return redirect(url_for('main.admin', site=site, request_deleted='1', **filters) + '#submissions')


def submission_dict(item):
    form = json.loads(item.form_json or '{}')
    result = json.loads(item.result_json or '{}')
    return {
        'id': item.id, 'site': item.site, 'form_kind': item.form_kind, 'name': item.name or '',
        'phone': item.phone or '', 'email': item.email or '',
        'employee': item.submitted_by.email if item.submitted_by else 'Public website',
        'details': form, 'details_display': display_fields(form),
        'result': result, 'result_display': display_fields(result), 'consent': 'Granted',
        'created_at': item.created_at.isoformat(), 'expires_at': item.expires_at.isoformat(),
        'quotation_pdf_url': (url_for('quotations.admin_pdf', lead_id=item.lead_id)
                              if item.site == 'batterywala' and item.form_kind == 'quotation' and item.lead_id else None),
    }


@bp.get('/admin/submissions.json')
@admin_required
def submissions_json():
    purge_expired_submissions()
    site, query = filtered_submissions(request.args)
    rows = query.limit(200).all()
    return jsonify(site=site, total=query.count(), rows=[submission_dict(item) for item in rows])


@bp.get('/admin/export.xlsx')
@admin_required
def export_submissions():
    purge_expired_submissions()
    site, query = filtered_submissions(request.args)
    selected = [name for name in request.args.getlist('columns') if name in EXPORT_COLUMNS]
    if not selected:
        selected = list(EXPORT_COLUMNS)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SITES[site][:31]
    sheet.append([EXPORT_COLUMNS[name] for name in selected])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='102E3C')
        cell.alignment = Alignment(vertical='center')
    for item in query.all():
        data = submission_dict(item)
        values = []
        for name in selected:
            value = data[name]
            if name == 'details':
                value = fields_as_text(data['details_display'])
            elif name == 'result':
                value = fields_as_text(data['result_display'])
            values.append(value)
        sheet.append(values)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(48, max(12, max(len(str(cell.value or '')) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
        for cell in column:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    stamp = utcnow().strftime('%Y-%m-%d')
    return send_file(output, as_attachment=True, download_name=f'{site}-submissions-{stamp}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
